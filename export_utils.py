# -*- coding: utf-8 -*-
"""
export_utils.py - 导出工具，兼容 Windows + Android
"""

import os
import csv
import sys
import datetime
import traceback

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    OPENPYXL_OK = True
except ImportError:
    OPENPYXL_OK = False

HEADERS = ["ID", "SSID", "BSSID", "信号(dBm)", "信道",
            "加密方式", "频段", "厂商", "采集时间", "关联单位"]
KEYS    = ["id", "ssid", "bssid", "signal", "channel",
           "encrypt_type", "band", "vendor", "collect_time", "unit_name"]


def _get_download_dir() -> str:
    """获取导出目录"""
    # Android：写到 /sdcard/Download
    ext = os.environ.get("EXTERNAL_STORAGE", "/sdcard")
    dldir = os.path.join(ext, "Download")
    if os.path.isdir(dldir):
        return dldir
    if os.path.isdir(ext):
        return ext
    return os.path.dirname(os.path.abspath(__file__))


def export_csv(rows: list, filepath: str) -> str:
    with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(HEADERS)
        for r in rows:
            w.writerow([r.get(k, "") for k in KEYS])
    return filepath


def export_excel(rows: list, filepath: str) -> str:
    if not OPENPYXL_OK:
        raise RuntimeError("openpyxl 未安装")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "WiFi数据"
    fill = PatternFill(fill_type="solid", fgColor="2F75B6")
    fnt  = Font(color="FFFFFF", bold=True)
    for ci, h in enumerate(HEADERS, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.fill = fill
        c.font = fnt
        c.alignment = Alignment(horizontal="center")
    for ri, r in enumerate(rows, 2):
        for ci, k in enumerate(KEYS, 1):
            ws.cell(row=ri, column=ci, value=r.get(k, ""))
    for col in ws.columns:
        mx = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(mx + 4, 40)
    wb.save(filepath)
    return filepath


def do_export(rows: list, fmt: str = "xlsx") -> tuple:
    """返回 (成功?, 文件路径或错误信息)"""
    ext = ".xlsx" if fmt == "xlsx" else ".csv"
    base = f"wifi_export_{datetime.datetime.now():%Y%m%d_%H%M%S}{ext}"
    ddir = _get_download_dir()
    filepath = os.path.join(ddir, base)
    try:
        if fmt == "xlsx":
            export_excel(rows, filepath)
        else:
            export_csv(rows, filepath)
        return True, filepath
    except Exception as e:
        traceback.print_exc()
        return False, str(e)
